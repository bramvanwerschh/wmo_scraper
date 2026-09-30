"""Conceptcodering: eerste-ronde, trefwoordgebaseerde suggestie per gemeente x
kernvraag uit het analysekader (Ecorys-offerte 1008762voo, Tabel 1.1). GEEN
inhoudelijke codering -- puur een geautomatiseerd voorstel dat Ingeborg/Sjoerd moeten
controleren en corrigeren (zelfde automatisering+handmatige-controle-model als de rest
van deze pijplijn, en zoals toegezegd in de offerte: "Onderzoekers controleren de
resultaten vervolgens handmatig op relevantie en coderen deze aan de hand van het
analysekader").

Hergebruikt de termenlijst en FILTER-logica uit src/termen.py (geen duplicatie, geen
risico op drift tussen de twee rapportages). "Kader" is geen kernvraag uit Tabel 1.1
maar wordt als 6e, informatieve rij meegenomen (de offerte noemt de handreiking
Gelijkgerichte Mantelzorgondersteuning expliciet als onderdeel van het analysekader).
"""

import csv
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

from src.config import ROOT
from src.termen import STARTLIJST, _in_filter, _matcht, _term_naar_patroon

csv.field_size_limit(sys.maxsize)

# Offerte Tabel 1.1 gebruikt deze exacte namen voor de 5 kernvragen; STARTLIJST-sleutels
# zijn de kortere werktitels die elders in dit project al gebruikt worden (o.a.
# termenverkenning.xlsx). Alleen het label voor de mensen die dit lezen verandert hier --
# de onderliggende termen/patronen blijven identiek aan src/termen.py.
KERNVRAAG_LABELS = {
    "Positie": "Positie van de mantelzorger",
    "Draagkracht": "Draagkracht en draaglast",
    "Behoeften": "Behoeften van de mantelzorger",
    "Instrumenten": "Instrumenten en werkwijzen",
    "Jonge mantelzorgers": "Jonge mantelzorgers",
    "Kader": "Kader (context, geen kernvraag uit Tabel 1.1)",
}

VOORBEELDEN_PER_GEMEENTE = 3


def genereer_conceptcodering(config: dict) -> dict:
    """Geeft {gemeente_code: {thema: {...}}} terug. Puur trefwoordmatching binnen de
    FILTER-scope (kern/waardering, domein wmo/beide, geen begrippen/bijlage/
    ondertekening/aanhef) -- geen semantisch begrip, geen garantie op volledigheid."""
    with open(ROOT / config["paden"]["passages_csv"], encoding="utf-8") as f:
        passages = list(csv.DictReader(f))
    with open(ROOT / config["paden"]["gemeenten_csv"], encoding="utf-8") as f:
        gemeenten = list(csv.DictReader(f))

    per_gemeente: dict[str, list[dict]] = {g["gemeente_code"]: [] for g in gemeenten}
    for p in passages:
        if p["gemeente_code"] in per_gemeente and _in_filter(p):
            per_gemeente[p["gemeente_code"]].append(p)

    thema_patronen = {
        thema: [(term, _term_naar_patroon(term)) for term in termen]
        for thema, termen in STARTLIJST.items()
    }

    resultaat: dict[str, dict] = {}
    for g in gemeenten:
        code = g["gemeente_code"]
        kandidaten = per_gemeente[code]
        thema_resultaat = {}
        for thema, term_patronen in thema_patronen.items():
            gematchte_termen = []
            treffers: list[dict] = []
            gezien_passage_ids = set()
            for term, patronen in term_patronen:
                term_treffers = [p for p in kandidaten if _matcht(p["tekst"], patronen)]
                if term_treffers:
                    gematchte_termen.append(term)
                    for t in term_treffers:
                        if t["passage_id"] not in gezien_passage_ids:
                            gezien_passage_ids.add(t["passage_id"])
                            treffers.append(t)

            thema_resultaat[thema] = {
                "voorgestelde_indicatie": "mogelijk aanwezig" if treffers else "niet aangetroffen in openbare documenten",
                "n_treffers": len(treffers),
                "n_termen_gematcht": len(gematchte_termen),
                "gematchte_termen": gematchte_termen,
                "voorbeelden": treffers[:VOORBEELDEN_PER_GEMEENTE],
            }
        resultaat[code] = thema_resultaat
    return resultaat


def _bronverwijzing(p: dict) -> str:
    onderdeel = p.get("artikel_titel") or p.get("hoofdstuk_titel") or p.get("pad") or ""
    return f"{p['cvdr_id']}/{p['versie']} — {onderdeel}".strip(" —")


def schrijf_codering_xlsx(config: dict, resultaat: dict):
    with open(ROOT / config["paden"]["gemeenten_csv"], encoding="utf-8") as f:
        gemeenten = sorted(csv.DictReader(f), key=lambda g: g["gemeente_naam_cbs"])

    wb = openpyxl.Workbook()

    ws_lees = wb.active
    ws_lees.title = "Leeswijzer"
    leeswijzer = [
        "CONCEPTCODERING -- eerste, geautomatiseerde suggestie per gemeente x kernvraag",
        "",
        "Dit is GEEN inhoudelijke codering. Elke 'voorgestelde_indicatie' komt puur uit "
        "trefwoordmatching binnen kern/waardering-passages met domein_hint wmo/beide "
        "(dezelfde FILTER-scope als termenverkenning.xlsx). Er is geen semantisch begrip: "
        "een treffer betekent dat een van de trefwoorden letterlijk voorkomt in een "
        "passage, niet dat de gemeente het onderwerp inhoudelijk goed heeft geregeld.",
        "",
        "'niet aangetroffen in openbare documenten' betekent dat geen van de trefwoorden "
        "een treffer gaf -- dit betekent NIET automatisch dat het onderwerp in de praktijk "
        "geen aandacht krijgt (zelfde nuance als de offerte zelf hanteert voor de "
        "documentanalyse-stap).",
        "",
        "Elk tabblad (per kernvraag) heeft twee lege kolommen 'definitieve_code' en "
        "'opmerking' -- vul die in tijdens de controle. Het tabblad 'Overzicht' geeft een "
        "landelijk beeld in één regel per gemeente, gebaseerd op dezelfde ongecontroleerde "
        "voorstellen.",
        "",
        "Brongegevens: data/passages.csv (kolom 'pad' geeft de structuurlocatie binnen het "
        "document; cvdr_id/versie zijn de identifiers om het brondocument op te zoeken in "
        "CVDR/lokaleregelgeving.overheid.nl).",
        "",
        "Termenlijst en FILTER-logica: src/termen.py (STARTLIJST) -- identiek aan "
        "termenverkenning.xlsx, dus die twee rapportages blijven per definitie consistent.",
    ]
    for i, regel in enumerate(leeswijzer, start=1):
        cel = ws_lees.cell(row=i, column=1, value=regel)
        if i == 1:
            cel.font = Font(bold=True, size=13)
        cel.alignment = Alignment(wrap_text=True)
    ws_lees.column_dimensions["A"].width = 110

    ws_overzicht = wb.create_sheet(title="Overzicht")
    kop = ["gemeente_naam", "gemeente_code"]
    for thema in STARTLIJST:
        kop += [f"{thema}: indicatie", f"{thema}: n_treffers"]
    ws_overzicht.append(kop)
    for g in gemeenten:
        code = g["gemeente_code"]
        rij = [g["gemeente_naam_cbs"], code]
        for thema in STARTLIJST:
            r = resultaat[code][thema]
            rij += [r["voorgestelde_indicatie"], r["n_treffers"]]
        ws_overzicht.append(rij)

    for thema in STARTLIJST:
        ws = wb.create_sheet(title=KERNVRAAG_LABELS[thema][:31])
        ws.append([
            "gemeente_naam", "gemeente_code",
            "voorgestelde_indicatie", "n_treffers", "n_termen_gematcht", "gematchte_termen",
            "voorbeeld_bron_1", "voorbeeld_tekst_1",
            "voorbeeld_bron_2", "voorbeeld_tekst_2",
            "voorbeeld_bron_3", "voorbeeld_tekst_3",
            "definitieve_code (in te vullen)", "opmerking (in te vullen)",
        ])
        for g in gemeenten:
            code = g["gemeente_code"]
            r = resultaat[code][thema]
            voorbeelden = r["voorbeelden"]
            rij = [
                g["gemeente_naam_cbs"], code,
                r["voorgestelde_indicatie"], r["n_treffers"], r["n_termen_gematcht"],
                "; ".join(r["gematchte_termen"]),
            ]
            for i in range(VOORBEELDEN_PER_GEMEENTE):
                if i < len(voorbeelden):
                    v = voorbeelden[i]
                    rij += [_bronverwijzing(v), v["tekst"][:400]]
                else:
                    rij += ["", ""]
            rij += ["", ""]
            ws.append(rij)
        ws.column_dimensions["A"].width = 22
        for col in ("H", "J", "L"):
            ws.column_dimensions[col].width = 60

    pad = ROOT / config["paden"]["rapportage"] / "codering_conceptscores.xlsx"
    pad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pad)
