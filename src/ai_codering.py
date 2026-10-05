"""AI-analyse (pilot): inhoudelijke conceptcodering per gemeente via Claude Sonnet 5.5
op Ecorys' Azure-omgeving (Microsoft Foundry).

Verschil met src/codering.py: dat is trefwoordmatching (aanwezig/niet aangetroffen,
geen begrip van INHOUD). Dit hier laat het model de gematchte passages per gemeente
ECHT lezen en per kernvraag een inhoudelijk oordeel geven: wat is er precies geregeld,
via welk documenttype (verordening/beleidsregel/...), en waarom (met brongegevens om
te verifieren) -- precies het onderscheid dat de Ecorys-offerte (1008762voo, Tabel 1.1)
vraagt: "niet alleen of een onderwerp wordt genoemd, maar ook op welke wijze dit is
uitgewerkt".

Nadrukkelijk een PILOT: bedoeld om eerst op een beperkte, al bekende steekproef
(de 10 A7-steekproefgemeenten) te draaien en door Ingeborg/Sjoerd te laten valideren
tegen de brontekst, voordat (op basis van die validatie) besloten wordt om op te
schalen naar alle 342 gemeenten. Kost echt geld (Ecorys' Azure-omgeving) -- zie
logs/werklog.md voor de kosteninschatting.
"""

import csv
import os
import sys
from pathlib import Path
from typing import Literal

import openpyxl
from dotenv import load_dotenv
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, Field

from anthropic import AnthropicFoundry

from src.config import ROOT
from src.termen import STARTLIJST, _in_filter, _matcht, _term_naar_patroon

csv.field_size_limit(sys.maxsize)

MODEL = "claude-sonnet-5-5"

KERNVRAAG_LABELS = {
    "Positie": "Positie van de mantelzorger",
    "Draagkracht": "Draagkracht en draaglast",
    "Behoeften": "Behoeften van de mantelzorger",
    "Instrumenten": "Instrumenten en werkwijzen",
    "Jonge mantelzorgers": "Jonge mantelzorgers",
    "Kader": "Kader (context, geen kernvraag uit Tabel 1.1)",
}

KERNVRAAG_OMSCHRIJVING = {
    "Positie": "Hoe en wanneer mantelzorgers worden betrokken bij het keukentafelgesprek, "
               "in hoeverre zij ruimte krijgen hun eigen situatie te bespreken en hoe hun "
               "perspectief meeweegt in de besluitvorming.",
    "Draagkracht": "Hoe de gemeente zicht krijgt op de belasting, belastbaarheid en "
                   "volhoudbaarheid van mantelzorg en welke factoren daarbij worden meegenomen.",
    "Behoeften": "Op welke wijze ondersteuningsbehoeften van mantelzorgers worden "
                 "geinventariseerd en hoe daar invulling aan wordt gegeven.",
    "Instrumenten": "Welke methodieken, hulpmiddelen en instrumenten worden voorgeschreven "
                     "om mantelzorgers te betrekken bij het keukentafelgesprek en de toegang "
                     "tot ondersteuning.",
    "Jonge mantelzorgers": "Hoe jonge mantelzorgers worden herkend, betrokken en ondersteund "
                            "binnen bovengenoemde thema's.",
    "Kader": "Of en hoe de gemeente verwijst naar landelijk kader (Mantelzorgagenda, "
             "handreiking Gelijkgerichte Mantelzorgondersteuning).",
}

STEEKPROEFGEMEENTEN_A7 = [
    "Hilvarenbeek", "Voerendaal", "Zoeterwoude", "Kaag en Braassem", "Amersfoort",
    "Voorst", "Bergen (L.)", "Geertruidenberg", "Barendrecht", "Leiden",
]


class KernvraagOordeel(BaseModel):
    indicatie: Literal["aanwezig", "gedeeltelijk aanwezig", "niet aangetroffen in openbare documenten"] = Field(
        description="Alleen baseren op de letterlijk gegeven brontekst, niet op algemene kennis."
    )
    documenttype: str = Field(
        description="Welk(e) documenttype(n) de vondst draagt, bijv. 'verordening', "
                     "'beleidsregel', 'verordening + beleidsregel', of 'n.v.t.' bij niet aangetroffen."
    )
    toelichting: str = Field(
        description="1-3 zinnen: WAT is er geregeld en HOE (concrete eis/instrument/norm), "
                     "niet alleen dat het onderwerp genoemd wordt. Nederlands, feitelijk, geen aannames."
    )
    bronnen: list[str] = Field(
        description="Lijst van gebruikte bronverwijzingen uit de aangeleverde passages "
                     "(cvdr_id/versie + artikel-aanduiding), leeg als niet aangetroffen."
    )


class GemeenteCodering(BaseModel):
    positie: KernvraagOordeel
    draagkracht: KernvraagOordeel
    behoeften: KernvraagOordeel
    instrumenten: KernvraagOordeel
    jonge_mantelzorgers: KernvraagOordeel
    kader: KernvraagOordeel


THEMA_NAAR_VELD = {
    "Positie": "positie",
    "Draagkracht": "draagkracht",
    "Behoeften": "behoeften",
    "Instrumenten": "instrumenten",
    "Jonge mantelzorgers": "jonge_mantelzorgers",
    "Kader": "kader",
}

SYSTEEM_PROMPT = """Je bent onderzoeksassistent voor Ecorys, in een onderzoek naar de positie \
van mantelzorgers bij de toegang tot Wmo-ondersteuning (opdracht VWS). Je krijgt per gemeente \
een verzameling brontekstpassages uit geldende Wmo-verordeningen/beleidsregels/nadere regels \
(CVDR), voorgeselecteerd op trefwoorden per kernvraag.

Beoordeel voor elke kernvraag uit het analysekader WAT er precies is geregeld en HOE -- niet \
alleen of het onderwerp genoemd wordt. Maak daarbij onderscheid tussen documenttypen (een \
verplichting in een verordening is iets anders dan een vrijblijvende vermelding in een \
beleidsregel of toelichting).

Regels:
- Baseer je oordeel UITSLUITEND op de letterlijk aangeleverde passages. Verzin niets en vul \
niets aan met algemene kennis over Wmo-beleid.
- 'niet aangetroffen in openbare documenten' betekent alleen dat de aangeleverde passages geen \
treffer bevatten -- dit is geen oordeel over de praktijk.
- Citeer bij elke 'aanwezig'/'gedeeltelijk aanwezig'-indicatie minstens 1 concrete bron \
(cvdr_id/versie + artikel-aanduiding) uit de aangeleverde tekst.
- Wees beknopt (1-3 zinnen per kernvraag)."""


def _gematchte_passages_per_thema(passages: list[dict]) -> dict[str, list[dict]]:
    thema_patronen = {
        thema: [(term, _term_naar_patroon(term)) for term in termen]
        for thema, termen in STARTLIJST.items()
    }
    kandidaten = [p for p in passages if _in_filter(p)]
    resultaat: dict[str, list[dict]] = {}
    for thema, term_patronen in thema_patronen.items():
        gezien = set()
        treffers = []
        for term, patronen in term_patronen:
            for p in kandidaten:
                if p["passage_id"] not in gezien and _matcht(p["tekst"], patronen):
                    gezien.add(p["passage_id"])
                    treffers.append(p)
        resultaat[thema] = treffers
    return resultaat


def _bronverwijzing(p: dict) -> str:
    onderdeel = p.get("artikel_titel") or p.get("hoofdstuk_titel") or p.get("pad") or ""
    return f"{p['cvdr_id']}/{p['versie']} — {onderdeel}".strip(" —")


def _bouw_prompt(gemeente_naam: str, per_thema: dict[str, list[dict]]) -> str:
    delen = [f"Gemeente: {gemeente_naam}\n"]
    for thema, treffers in per_thema.items():
        delen.append(f"\n=== Kernvraag: {KERNVRAAG_LABELS[thema]} ===")
        delen.append(f"({KERNVRAAG_OMSCHRIJVING[thema]})")
        if not treffers:
            delen.append("(geen trefwoordtreffers voor deze kernvraag)")
            continue
        for p in treffers:
            delen.append(f"\n[{_bronverwijzing(p)}]\n{p['tekst']}")
    delen.append(
        "\n\nGeef voor elke kernvraag hierboven een oordeel volgens het gevraagde "
        "outputformat (velden: positie, draagkracht, behoeften, instrumenten, "
        "jonge_mantelzorgers, kader)."
    )
    return "\n".join(delen)


def _client() -> AnthropicFoundry:
    load_dotenv(ROOT / ".env")
    endpoint = os.environ["AZURE_CLAUDE_ENDPOINT"]
    resource = endpoint.split("//", 1)[1].split(".", 1)[0]
    return AnthropicFoundry(api_key=os.environ["AZURE_CLAUDE_API_KEY"], resource=resource)


def codeer_gemeente(client: AnthropicFoundry, gemeente_naam: str, passages: list[dict]) -> GemeenteCodering:
    per_thema = _gematchte_passages_per_thema(passages)
    prompt = _bouw_prompt(gemeente_naam, per_thema)
    response = client.messages.parse(
        model=MODEL,
        max_tokens=4096,
        system=SYSTEEM_PROMPT,
        messages=[{"role": "user", "content": prompt}],
        output_format=GemeenteCodering,
    )
    return response.parsed_output


def draai_pilot(config: dict, gemeenten_namen: list[str] | None = None) -> dict[str, GemeenteCodering]:
    """Draait de AI-codering voor een beperkte lijst gemeenten (default: de 10
    A7-steekproefgemeenten). Kost echte API-aanroepen/geld -- niet voor de volledige
    342 gemeenten zonder expliciet besluit na validatie van deze pilot."""
    gemeenten_namen = gemeenten_namen or STEEKPROEFGEMEENTEN_A7
    resultaat, fouten = _draai(config, gemeenten_namen)
    if fouten:
        raise RuntimeError(f"Pilot-fouten bij: {fouten}")
    return resultaat


def draai_alle_gemeenten(config: dict) -> tuple[dict[str, GemeenteCodering], list[str]]:
    """Volledige run over alle 342 gemeenten. Kost echte API-aanroepen/geld (~EUR 23-25,
    Sonnet 5.5, ~60-65 min). Per-gemeente foutafhandeling: een mislukte aanroep stopt de
    run niet, de naam komt in de teruggegeven foutenlijst terecht (logging per 10
    gemeenten voor voortgang bij een lange achtergrondrun)."""
    with open(ROOT / config["paden"]["gemeenten_csv"], encoding="utf-8") as f:
        alle_namen = [g["gemeente_naam_cbs"] for g in csv.DictReader(f)]
    return _draai(config, alle_namen, log_voortgang=True)


CACHE_DIR = "data/raw/ai_codering_cache"


def _cache_pad(code: str) -> Path:
    return ROOT / CACHE_DIR / f"{code}.json"


def _laad_cache(code: str) -> GemeenteCodering | None:
    pad = _cache_pad(code)
    if not pad.exists():
        return None
    return GemeenteCodering.model_validate_json(pad.read_text(encoding="utf-8"))


def _schrijf_cache(code: str, codering: GemeenteCodering):
    pad = _cache_pad(code)
    pad.parent.mkdir(parents=True, exist_ok=True)
    pad.write_text(codering.model_dump_json(), encoding="utf-8")


def _draai(
    config: dict, gemeenten_namen: list[str], log_voortgang: bool = False
) -> tuple[dict[str, GemeenteCodering], list[str]]:
    """Hervatbaar: elk gemeente-resultaat wordt direct na de API-aanroep weggeschreven
    naar data/raw/ai_codering_cache/<code>.json (zelfde cache-filosofie als de rest van
    de pijplijn). Bij een onderbroken run (crash, tijdslimiet, laptop dicht) kost een
    herstart alleen nieuwe API-aanroepen voor de gemeenten die nog geen cachebestand
    hebben -- geen dubbel werk, geen dubbele kosten."""
    with open(ROOT / config["paden"]["gemeenten_csv"], encoding="utf-8") as f:
        gemeenten = list(csv.DictReader(f))
    naam_naar_code = {g["gemeente_naam_cbs"]: g["gemeente_code"] for g in gemeenten}

    with open(ROOT / config["paden"]["passages_csv"], encoding="utf-8") as f:
        alle_passages = list(csv.DictReader(f))

    client = _client()
    resultaat: dict[str, GemeenteCodering] = {}
    fouten: list[str] = []
    n_uit_cache = 0
    for i, naam in enumerate(gemeenten_namen, start=1):
        code = naam_naar_code[naam]
        cache_treffer = _laad_cache(code)
        if cache_treffer is not None:
            resultaat[naam] = cache_treffer
            n_uit_cache += 1
        else:
            passages = [p for p in alle_passages if p["gemeente_code"] == code]
            try:
                codering = codeer_gemeente(client, naam, passages)
                _schrijf_cache(code, codering)
                resultaat[naam] = codering
            except Exception as e:
                fouten.append(naam)
                if log_voortgang:
                    print(f"  FOUT bij {naam}: {e}", flush=True)
        if log_voortgang and i % 10 == 0:
            print(f"  {i}/{len(gemeenten_namen)} verwerkt ({n_uit_cache} uit cache, {len(fouten)} fouten)", flush=True)
    return resultaat, fouten


def schrijf_pilot_xlsx(
    config: dict, resultaat: dict[str, GemeenteCodering], pad: Path | None = None,
    volledig: bool = False, fouten: list[str] | None = None,
):
    wb = openpyxl.Workbook()
    ws_lees = wb.active
    ws_lees.title = "Leeswijzer"
    if volledig:
        titel = "AI-CODERING -- ALLE GEMEENTEN (Claude Sonnet 5.5, Ecorys Azure/Foundry)"
        scope_regel = (
            f"DIT IS DE VOLLEDIGE RUN (alle {len(resultaat)} gemeenten), op besluit van Bram "
            "(2026-10-05) om door te gaan zodat Ingeborg/Sjoerd hiermee aan de slag kunnen, "
            "ook al is de pilot-validatie door een onafhankelijke lezer nog niet afgerond. "
            "Claude's eigen steekproefcontrole (20 rijen van de 10-gemeenten-pilot, tegen de "
            "volledige brontekst) gaf 0 fouten -- geen vervanging voor een onafhankelijke "
            "lezer, wel een reden om door te gaan."
        )
        if fouten:
            scope_regel += f" MISLUKT voor {len(fouten)} gemeente(n): {', '.join(fouten)}."
    else:
        titel = "AI-CODERING PILOT (Claude Sonnet 5.5, Ecorys Azure/Foundry)"
        scope_regel = (
            "DIT IS EEN PILOT OP 10 GEMEENTEN. Voordat dit wordt opgeschaald naar alle 342 "
            "gemeenten, moet elke rij hier gecontroleerd worden tegen de brontekst (kolommen "
            "'gevalideerd_door' en 'opmerking_validatie' invullen). Bij systematische fouten: "
            "prompt aanpassen en pilot herhalen voordat wordt opgeschaald."
        )
    leeswijzer = [
        titel,
        "",
        "Dit IS een inhoudelijke beoordeling (i.t.t. codering_conceptscores.xlsx, dat pure "
        "trefwoordmatching is). Het model heeft de gematchte brontekst per gemeente echt "
        "gelezen en per kernvraag een oordeel gegeven: wat is geregeld, via welk "
        "documenttype, met brongegevens erbij.",
        "",
        scope_regel,
        "",
        "Brongegevens: data/passages.csv. Zelfde termenlijst/FILTER-scope als "
        "termenverkenning.xlsx en codering_conceptscores.xlsx (src/termen.py).",
    ]
    for i, regel in enumerate(leeswijzer, start=1):
        cel = ws_lees.cell(row=i, column=1, value=regel)
        if i == 1:
            cel.font = Font(bold=True, size=13)
        cel.alignment = Alignment(wrap_text=True)
    ws_lees.column_dimensions["A"].width = 110

    ws = wb.create_sheet(title="Pilot-resultaten")
    ws.append([
        "gemeente_naam", "kernvraag", "indicatie", "documenttype", "toelichting", "bronnen",
        "gevalideerd_door (in te vullen)", "opmerking_validatie (in te vullen)",
    ])
    for gemeente_naam, codering in resultaat.items():
        for thema, veld in THEMA_NAAR_VELD.items():
            oordeel: KernvraagOordeel = getattr(codering, veld)
            ws.append([
                gemeente_naam, KERNVRAAG_LABELS[thema], oordeel.indicatie, oordeel.documenttype,
                oordeel.toelichting, "; ".join(oordeel.bronnen), "", "",
            ])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["E"].width = 70
    ws.column_dimensions["F"].width = 40

    pad = pad or (ROOT / config["paden"]["rapportage"] / "ai_codering_pilot.xlsx")
    pad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pad)
